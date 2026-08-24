"""
Fast, standalone refresh of just the "Payment Plans" sheet in the Zisla
inventory tracker workbook. Unlike track_zisla_inventory.py (which also
walks every unit in every development and takes ~10+ minutes), this only
needs the development listing, so it runs in seconds - handy for an
on-demand refresh between monthly full runs.

Payment plan data comes from each development's "financing" field: one or
more plan options, each a % breakdown due at signing / during construction /
at delivery / at deeding, plus any discount for that plan.

Appends a dated snapshot to the SAME workbook used by track_zisla_inventory.py
(default: G:\\My Drive\\zisla-tracker\\zisla-inventory-tracker.xlsx), touching
only the "Payment Plans" sheet - Units/Developments are left untouched.

Usage:
  python download_zisla_payment_plans.py
  python download_zisla_payment_plans.py --out "D:\\other\\path.xlsx"
"""
from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import load_workbook

from zisla_common import SNAPSHOT_DATE, append_rows, fetch_developments, get_or_create_sheet


def build_plan_rows(developments: list[dict]) -> list[dict]:
    rows = []
    for p in developments:
        financing = p.get("financing") or []
        address = p.get("address") or {}
        localized = p.get("localized") or {}
        for plan_num, plan in enumerate(financing, start=1):
            rows.append(
                {
                    "Snapshot Date": SNAPSHOT_DATE,
                    "Development": localized.get("title"),
                    "Development ID": p.get("_id"),
                    "Province": address.get("province"),
                    "City": address.get("city"),
                    "Plan Number": plan_num,
                    "Plan ID": plan.get("_id"),
                    "Signin %": plan.get("signin"),
                    "Building %": plan.get("building"),
                    "Delivery %": plan.get("delivery"),
                    "Deeding %": plan.get("deeding"),
                    "Discount %": plan.get("discount"),
                }
            )
    return rows


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--out", default=r"G:\My Drive\zisla-tracker\zisla-inventory-tracker.xlsx")
    args = parser.parse_args()

    print("Fetching developments from zisla.com ...")
    developments = fetch_developments()
    print(f"Total developments: {len(developments)}")

    print("Building payment plan rows...")
    plan_rows = build_plan_rows(developments)
    print(f"Total payment plan rows: {len(plan_rows)}")

    if not plan_rows:
        print("No payment plans found - nothing to write.")
        return

    out_path = Path(args.out)
    if not out_path.exists():
        raise SystemExit(f"Workbook not found at {out_path} - run track_zisla_inventory.py first to create it.")

    print(f"Writing to workbook: {out_path}")
    wb = load_workbook(out_path)
    ws = get_or_create_sheet(wb, "Payment Plans", list(plan_rows[0].keys()))

    print(f"Appending {len(plan_rows)} payment plan rows...")
    append_rows(ws, plan_rows)

    total_rows = ws.max_row - 1
    wb.save(out_path)
    print(f"Saved snapshot dated {SNAPSHOT_DATE} to {out_path}")
    print(f"Payment Plans sheet now has {total_rows} total rows across all snapshots")


if __name__ == "__main__":
    main()
