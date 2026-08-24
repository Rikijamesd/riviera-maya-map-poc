"""
Builds a Word document listing every published Zisla blog article title,
grouped by category - a quick-scan companion to the data in
download_zisla_blog_titles.py's spreadsheet output.

Usage:
  python make_blog_titles_docx.py
  python make_blog_titles_docx.py "D:\\other\\zisla-blog-titles.docx"
"""
from __future__ import annotations

import sys
from datetime import date

from docx import Document
from docx.shared import Pt, RGBColor

from zisla_common import fetch_blog_articles

CATEGORY_ORDER = ["TIPS", "NEWS", "NEW_PROJECTS", "BUILDER_DESIGN", "FINANCES", "LIFESTYLE", "INVEST", "TOURISM"]
CATEGORY_LABELS = {
    "TIPS": "Tips",
    "NEWS": "News",
    "NEW_PROJECTS": "New Projects",
    "BUILDER_DESIGN": "Builder Design",
    "FINANCES": "Finances",
    "LIFESTYLE": "Lifestyle",
    "INVEST": "Investing",
    "TOURISM": "Tourism",
}


def main() -> None:
    out_path = sys.argv[1] if len(sys.argv) > 1 else "zisla-blog-titles.docx"

    print("Fetching blog articles...")
    articles = fetch_blog_articles()
    print(f"Fetched {len(articles)} articles")

    by_category: dict[str, list[str]] = {}
    for a in articles:
        cat = a.get("blogType") or "UNCATEGORIZED"
        title = ((a.get("localized") or {}).get("en") or {}).get("title")
        by_category.setdefault(cat, []).append(title)

    ordered_cats = [c for c in CATEGORY_ORDER if c in by_category]
    ordered_cats += [c for c in by_category if c not in CATEGORY_ORDER]

    doc = Document()
    doc.add_heading("Zisla Blog Article Titles", level=1)

    subtitle = doc.add_paragraph()
    run = subtitle.add_run(f"{len(articles)} published articles, scraped {date.today().isoformat()}.")
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    for cat in ordered_cats:
        titles = by_category[cat]
        doc.add_heading(f"{CATEGORY_LABELS.get(cat, cat)} ({len(titles)})", level=2)
        for title in titles:
            doc.add_paragraph(title)

    doc.save(out_path)
    print(f"Saved {len(articles)} article titles across {len(ordered_cats)} categories to {out_path}")


if __name__ == "__main__":
    main()
