"""
Fast, standalone refresh of just the "Phases" sheet in the Zisla inventory
tracker workbook. Only needs the development listing (no per-unit fetch), so
it runs in seconds - handy for an on-demand refresh between monthly full runs.

Phase/delivery data comes from each development's "statusByPhase" field: one
or more construction phases, each with a status (e.g. AVAILABLE,
READY_TO_MOVE_IN) and a delivery date.

Appends a dated snapshot to the SAME workbook used by track_zisla_inventory.py
(default: G:\\My Drive\\zisla-tracker\\zisla-inventory-tracker.xlsx), touching
only the "Phases" sheet - Units/Developments/Payment Plans are left untouched.

Usage:
  python download_zisla_phases.py
  python download_zisla_phases.py --out "D:\\other\\path.xlsx"
"""
from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import load_workbook

from zisla_common import SNAPSHOT_DATE, append_rows, fetch_developments, get_or_create_sheet


def build_phase_rows(developments: list[dict]) -> list[dict]:
    rows = []
    for p in developments:
        phases = p.get("statusByPhase") or []
        address = p.get("address") or {}
        localized = p.get("localized") or {}
        for phase_num, phase in enumerate(phases, start=1):
            rows.append(
                {
                    "Snapshot Date": SNAPSHOT_DATE,
                    "Development": localized.get("title"),
                    "Development ID": p.get("_id"),
                    "Province": address.get("province"),
                    "City": address.get("city"),
                    "Phase Number": phase_num,
                    "Phase ID": phase.get("_id"),
                    "Status": phase.get("status"),
                    "Delivery Date": phase.get("deliveryDate"),
                }
            )
    return rows


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--out", default=r"G:\My Drive\zisla-tracker\zisla-inventory-tracker.xlsx")
    args = parser.parse_args()

    print("Fetching developments from zisla.com ...")
    developments = fetch_developments()
    print(f"Total developments: {len(developments)}")

    print("Building phase/delivery rows...")
    phase_rows = build_phase_rows(developments)
    print(f"Total phase rows: {len(phase_rows)}")

    if not phase_rows:
        print("No phase data found - nothing to write.")
        return

    out_path = Path(args.out)
    if not out_path.exists():
        raise SystemExit(f"Workbook not found at {out_path} - run track_zisla_inventory.py first to create it.")

    print(f"Writing to workbook: {out_path}")
    wb = load_workbook(out_path)
    ws = get_or_create_sheet(wb, "Phases", list(phase_rows[0].keys()))

    print(f"Appending {len(phase_rows)} phase rows...")
    append_rows(ws, phase_rows)

    total_rows = ws.max_row - 1
    wb.save(out_path)
    print(f"Saved snapshot dated {SNAPSHOT_DATE} to {out_path}")
    print(f"Phases sheet now has {total_rows} total rows across all snapshots")


if __name__ == "__main__":
    main()
