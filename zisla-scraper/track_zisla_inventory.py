"""
Monthly inventory tracker for Zisla (zisla.com).

Fetches every published development plus every individual unit within it,
and APPENDS a dated snapshot to a persistent Excel workbook (rather than
overwriting), so you can see how availability, pricing, and inventory
change over time.

Sheets:
  - Units         : one row per unit per snapshot (Development + Unit columns)
  - Developments  : one row per development per snapshot (project-level info)
  - Payment Plans : one row per financing/payment plan option per development per snapshot
  - Phases        : one row per construction phase per development per snapshot (status + delivery date)

Default output: G:\\My Drive\\zisla-tracker\\zisla-inventory-tracker.xlsx (synced via
Google Drive for Desktop).

Usage:
  python track_zisla_inventory.py
  python track_zisla_inventory.py --out "D:\\other\\path.xlsx"
  python track_zisla_inventory.py --test          # only first 5 developments, for a quick dry run
"""
from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import Workbook, load_workbook

from zisla_common import (
    SNAPSHOT_DATE,
    append_rows,
    autofit_columns,
    fetch_developments,
    fetch_units,
    get_or_create_sheet,
    join_true_keys,
    to_plain_text,
)


def build_dev_rows(developments: list[dict]) -> list[dict]:
    rows = []
    for p in developments:
        price = p.get("price") or {}
        address = p.get("address") or {}
        localized = p.get("localized") or {}
        rows.append(
            {
                "Snapshot Date": SNAPSHOT_DATE,
                "Development": localized.get("title"),
                "Development ID": p.get("_id"),
                "URL": f"https://www.zisla.com/en/properties/{localized.get('urlAlias')}",
                "Purpose": p.get("purpose"),
                "Type": p.get("type"),
                "Sub Type": p.get("subType"),
                "Province": address.get("province"),
                "City": address.get("city"),
                "Price Min (USD)": round(price["min"] / 100, 2) if price.get("min") is not None else None,
                "Price Max (USD)": round(price["max"] / 100, 2) if price.get("max") is not None else None,
                "Units Total": p.get("nbOfUnits"),
                "Units Available": p.get("nbOfUnitsAvailable"),
                "Developer ID": p.get("developer"),
                "Financing Offered": p.get("isFinancing"),
                "Amenities": join_true_keys(p.get("amenities")),
                "Features": join_true_keys(p.get("features")),
                "Description": to_plain_text(localized.get("description")),
                "Last Update": p.get("lastUpdate"),
            }
        )
    return rows


def build_plan_rows(developments: list[dict]) -> list[dict]:
    rows = []
    for p in developments:
        financing = p.get("financing") or []
        address = p.get("address") or {}
        localized = p.get("localized") or {}
        for plan_num, plan in enumerate(financing, start=1):
            rows.append(
                {
                    "Snapshot Date": SNAPSHOT_DATE,
                    "Development": localized.get("title"),
                    "Development ID": p.get("_id"),
                    "Province": address.get("province"),
                    "City": address.get("city"),
                    "Plan Number": plan_num,
                    "Plan ID": plan.get("_id"),
                    "Signin %": plan.get("signin"),
                    "Building %": plan.get("building"),
                    "Delivery %": plan.get("delivery"),
                    "Deeding %": plan.get("deeding"),
                    "Discount %": plan.get("discount"),
                }
            )
    return rows


def build_phase_rows(developments: list[dict]) -> list[dict]:
    rows = []
    for p in developments:
        phases = p.get("statusByPhase") or []
        address = p.get("address") or {}
        localized = p.get("localized") or {}
        for phase_num, phase in enumerate(phases, start=1):
            rows.append(
                {
                    "Snapshot Date": SNAPSHOT_DATE,
                    "Development": localized.get("title"),
                    "Development ID": p.get("_id"),
                    "Province": address.get("province"),
                    "City": address.get("city"),
                    "Phase Number": phase_num,
                    "Phase ID": phase.get("_id"),
                    "Status": phase.get("status"),
                    "Delivery Date": phase.get("deliveryDate"),
                }
            )
    return rows


def build_unit_rows(developments: list[dict]) -> list[dict]:
    rows = []
    total = len(developments)
    for i, p in enumerate(developments, start=1):
        address = p.get("address") or {}
        localized = p.get("localized") or {}
        units = fetch_units(p["_id"])
        for u in units:
            phase = u.get("phase") or {}
            rows.append(
                {
                    "Snapshot Date": SNAPSHOT_DATE,
                    "Development": localized.get("title"),
                    "Development ID": p.get("_id"),
                    "Province": address.get("province"),
                    "City": address.get("city"),
                    "Unit": u.get("name"),
                    "Unit ID": u.get("_id"),
                    "Status": u.get("status"),
                    "Bedrooms": u.get("bedrooms"),
                    "Bathrooms": u.get("bathrooms"),
                    "Construct Area (m2)": u.get("squareSurface"),
                    "Price (USD)": round(u["price"] / 100, 2) if u.get("price") is not None else None,
                    "Price (MXN)": round(u["priceInPesos"] / 100, 2) if u.get("priceInPesos") is not None else None,
                    "Price per Area": u.get("priceSurface"),
                    "Floor Number": u.get("floorNumber"),
                    "Building ID": u.get("buildingId"),
                    "Phase Status": phase.get("status"),
                    "Phase Delivery Date": phase.get("deliveryDate"),
                    "Discounted": u.get("discount"),
                    "Unit Last Updated": u.get("updatedAt"),
                }
            )
        if i % 25 == 0 or i == total:
            print(f"  {i} / {total} developments processed, {len(rows)} units collected so far")
    return rows


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--out", default=r"G:\My Drive\zisla-tracker\zisla-inventory-tracker.xlsx")
    parser.add_argument("--test", action="store_true")
    args = parser.parse_args()

    print("Fetching developments from zisla.com ...")
    developments = fetch_developments()
    print(f"Total developments: {len(developments)}")

    if args.test:
        print("TEST MODE: limiting to first 5 developments")
        developments = developments[:5]

    print("Building development rows...")
    dev_rows = build_dev_rows(developments)

    print("Building payment plan rows...")
    plan_rows = build_plan_rows(developments)
    print(f"Total payment plan rows: {len(plan_rows)}")

    print("Building phase/delivery rows...")
    phase_rows = build_phase_rows(developments)
    print(f"Total phase rows: {len(phase_rows)}")

    print("Fetching units for each development (this takes a while)...")
    unit_rows = build_unit_rows(developments)
    print(f"Total unit rows: {len(unit_rows)}")

    out_path = Path(args.out)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    print(f"Writing to workbook: {out_path}")
    if out_path.exists():
        wb = load_workbook(out_path)
    else:
        wb = Workbook()
        wb.remove(wb.active)  # drop the blank default sheet - named sheets are added below

    units_sheet = get_or_create_sheet(wb, "Units", list(unit_rows[0].keys()))
    dev_sheet = get_or_create_sheet(wb, "Developments", list(dev_rows[0].keys()))
    plan_sheet = get_or_create_sheet(wb, "Payment Plans", list(plan_rows[0].keys())) if plan_rows else None
    phase_sheet = get_or_create_sheet(wb, "Phases", list(phase_rows[0].keys())) if phase_rows else None

    print(f"Appending {len(unit_rows)} unit rows...")
    append_rows(units_sheet, unit_rows)
    print(f"Appending {len(dev_rows)} development rows...")
    append_rows(dev_sheet, dev_rows)
    if plan_sheet is not None:
        print(f"Appending {len(plan_rows)} payment plan rows...")
        append_rows(plan_sheet, plan_rows)
    if phase_sheet is not None:
        print(f"Appending {len(phase_rows)} phase rows...")
        append_rows(phase_sheet, phase_rows)

    for ws in (units_sheet, dev_sheet, plan_sheet, phase_sheet):
        if ws is not None:
            autofit_columns(ws)

    units_total = units_sheet.max_row - 1
    dev_total = dev_sheet.max_row - 1
    plan_total = plan_sheet.max_row - 1 if plan_sheet is not None else 0
    phase_total = phase_sheet.max_row - 1 if phase_sheet is not None else 0

    wb.save(out_path)
    print(f"Saved snapshot dated {SNAPSHOT_DATE} to {out_path}")
    print(f"Units sheet now has {units_total} total rows across all snapshots")
    print(f"Developments sheet now has {dev_total} total rows across all snapshots")
    print(f"Payment Plans sheet now has {plan_total} total rows across all snapshots")
    print(f"Phases sheet now has {phase_total} total rows across all snapshots")


if __name__ == "__main__":
    main()
